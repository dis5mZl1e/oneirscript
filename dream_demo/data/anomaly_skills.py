"""异象主动技能数据定义层（设计提案 v1.0 · 2026-10-08）。

数据源：`梦层异象主动技能开发表.xlsx`（仓库根目录，openpyxl 只读解析）。
本模块在导入时一次性读取《技能参数》《技能效果》《异象配置》《演出预设》
《调参规则》五张工作表，并合入 hospital.py 的专属位面成品参数：

- ``SKILL_DEFINITIONS``：43 个主动技能定义（稳定技能 ID 为唯一键，废弃旧
  ``s0/s1`` 下标语义）。
- ``ANOMALY_SKILL_BINDINGS``：56 种异象的技能绑定，共 87 个绑定槽位。
- ``FX_PROFILES``：10 种局部粒子演出预设。
- 调参常量：四层段伤害倍率、各阶级伤害上限、共享起手间隔、首次警戒延迟、
  场/控制/召唤/治疗/演出预算等。

绑定覆写（弱化/教学/阶段）在 ``_BINDING_OVERRIDES`` 与 ``_CORE_PHASES``
中以结构化字典维护，与 xlsx《异象配置》的「绑定覆写/阶段」文本一一对应；
``override_note`` 保留 xlsx 原文便于核对。生效值用
``effective_skill_def()`` / ``bindings_for()`` 查询。

导入时执行 ``validate_or_raise()``：原表43/56/87加专属定义的数量、未知技能 ID、槽位
数量阶级规则（普通 0–2 / 强烈 2–3 / 核心 4）、CD/蓄力/恢复非负、
伤害上限阶层合法等任一项失败都会抛 ``RuntimeError``，避免脏数据静默进入。
"""

from __future__ import annotations

import json
import math
from dataclasses import dataclass, field, replace
from pathlib import Path
from types import MappingProxyType
from typing import Any, Mapping

_WORKBOOK_PATH = Path(__file__).resolve().parents[2] / "梦层异象主动技能开发表.xlsx"
# 预转换产物（tools/export_anomaly_skills_json.py 生成）：存在时导入期只读
# 本文件，不依赖 openpyxl 与 xlsx（网页端 Pyodide 构建的前提）。
_PREBUILT_PATH = Path(__file__).resolve().parent / "anomaly_skills.json"

_TIER_NAME_TO_INT = {"普通": 0, "强烈": 1, "核心": 2}
TIER_NAMES = ("普通", "强烈", "核心")

_LAYER_BAND_RANGES = ((1, 10), (11, 20), (21, 27), (28, 30))


# ---------------------------------------------------------------------------
# 数据类
# ---------------------------------------------------------------------------

@dataclass(frozen=True)
class SkillDef:
    """主动技能基础定义（对应《技能参数》+《技能效果》两行合并）。"""

    skill_id: str
    name_zh: str
    family: str                       # 类型：弹道/位移/负面征兆/效果场/…
    default_tier: int                 # 默认阶级 0 普通 / 1 强烈 / 2 核心
    cooldown_rounds: int              # CD：完成施放后到再次允许蓄力的客观回合
    windup_rounds: int                # 蓄力回合（W）
    max_range: int                    # 施法距离（切比雪夫）
    shape: str                        # 形状：固定细线/固定十字/…
    max_cells: int                    # 格数上限（坐标列表是权威，格数是上限）
    direct_base: int                  # 基础直接创伤（0 表示无直接创伤）
    damage_by_band: tuple[int, int, int, int]  # 四层段预计算直接创伤
    trauma_channel: str               # 创伤渠道：心智 / 清醒 / 无
    duration_rounds: int              # 场/征兆持续（客观回合）
    recovery_rounds: int              # 恢复回合（结算后不能普攻/起手）
    fx_profile: str                   # 演出预设 ID
    effect_text: str                  # 完整效果与叠加限制（图鉴详情）
    condition_text: str               # 合法施放条件
    counter_text: str                 # 玩家反制
    status: str                       # 开发状态（提案值，非已实装）
    fixed_damage: bool = False         # Authored exclusive-realm damage, already region-scaled.


@dataclass(frozen=True)
class SkillBinding:
    """一个异象的一个主动技能绑定槽位（对应《异象配置》一列主动 N）。"""

    anomaly_id: str
    skill_id: str
    slot: int                         # 配置槽顺序（0 起，即优先级）
    priority: int                     # 同优先级用「上次施放更早 → 槽顺序」
    phase: int                        # 核心阶段解锁（1 起）；非阶段绑定为 1
    overrides: Mapping[str, Any] = field(default_factory=dict)
    override_note: str = ""           # xlsx「绑定覆写/阶段」原文

    def __post_init__(self) -> None:
        object.__setattr__(self, "overrides", MappingProxyType(dict(self.overrides)))

    @property
    def effective_tier(self) -> int:
        """结算伤害上限时使用的阶级（默认取技能定义，可被覆写）。"""
        tier = self.overrides.get("effective_tier")
        return self.default_tier if tier is None else int(tier)

    @property
    def default_tier(self) -> int:
        return SKILL_DEFINITIONS[self.skill_id].default_tier


@dataclass(frozen=True)
class FxProfile:
    """局部粒子演出预设（对应《演出预设》）。"""

    profile_id: str
    color_role: str
    color_hex: str
    bright_hex: str
    particle_char: str
    warn_cells_high: int
    warn_cells_low: int
    cast_cells_high: int
    cast_cells_low: int
    entry_seconds: float
    cast_seconds: float
    tail_seconds: float
    range_hint: str


@dataclass(frozen=True)
class AnomalySkillConfig:
    """异象级主动技能配置（对应《异象配置》行级字段）。"""

    anomaly_id: str
    name_zh: str
    tier: int
    layer_start: int
    layer_end: int
    initial_delay_rounds: int         # 初次警戒延迟
    global_gap_rounds: int            # 共享起手间隔
    override_note: str
    migration_note: str
    status: str


# ---------------------------------------------------------------------------
# xlsx 解析
# ---------------------------------------------------------------------------

def _to_int(value: Any) -> int:
    if value is None or value == "":
        return 0
    if isinstance(value, (int, float)):
        return int(value)
    text = str(value).strip()
    return int(float(text)) if text else 0


def _to_float(value: Any) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    return float(str(value).strip() or 0.0)


def _to_str(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _rows_by_header(ws: Any, header_first_cell: str) -> list[dict[str, Any]]:
    """定位表头行（首个单元格匹配），返回其下所有非空首列行字典。"""
    header: list[str] | None = None
    rows: list[dict[str, Any]] = []
    for row in ws.iter_rows(values_only=True):
        cells = list(row)
        if header is None:
            if cells and _to_str(cells[0]) == header_first_cell:
                header = [_to_str(c) for c in cells]
            continue
        if not cells or not _to_str(cells[0]):
            continue
        rows.append({h: c for h, c in zip(header, cells) if h})
    return rows


def _load_workbook() -> Any:
    import openpyxl  # 懒导入：有预转换 JSON 时运行期不需要该依赖

    if not _WORKBOOK_PATH.exists():
        raise RuntimeError(f"异象主动技能开发表不存在：{_WORKBOOK_PATH}")
    return openpyxl.load_workbook(_WORKBOOK_PATH, read_only=True, data_only=True)


def _parse_skills(wb: openpyxl.Workbook) -> dict[str, SkillDef]:
    params = {r["技能ID"]: r for r in _rows_by_header(wb["技能参数"], "技能ID")}
    effects = {r["技能ID"]: r for r in _rows_by_header(wb["技能效果"], "技能ID")}
    definitions: dict[str, SkillDef] = {}
    for skill_id, p in params.items():
        e = effects.get(skill_id, {})
        tier_name = _to_str(p.get("默认阶级"))
        if tier_name not in _TIER_NAME_TO_INT:
            raise RuntimeError(f"技能 {skill_id} 默认阶级非法：{tier_name!r}")
        definitions[skill_id] = SkillDef(
            skill_id=skill_id,
            name_zh=_to_str(p.get("中文名")),
            family=_to_str(p.get("类型")),
            default_tier=_TIER_NAME_TO_INT[tier_name],
            cooldown_rounds=_to_int(p.get("CD回合")),
            windup_rounds=_to_int(p.get("蓄力回合")),
            max_range=_to_int(p.get("施法距离")),
            shape=_to_str(p.get("形状")),
            max_cells=_to_int(p.get("格数上限")),
            direct_base=_to_int(p.get("基础直接创伤")),
            damage_by_band=tuple(
                _to_int(p.get(f"{lo}–{hi}层"))
                for lo, hi in _LAYER_BAND_RANGES
            ),
            trauma_channel=_to_str(p.get("创伤渠道")),
            duration_rounds=_to_int(p.get("场/征兆持续")),
            recovery_rounds=_to_int(p.get("恢复回合")),
            fx_profile=_to_str(p.get("演出预设")),
            effect_text=_to_str(e.get("完整效果与叠加限制")),
            condition_text=_to_str(e.get("合法施放条件")),
            counter_text=_to_str(e.get("玩家反制")),
            status=_to_str(p.get("开发状态")),
        )
    return definitions


def _parse_fx_profiles(wb: openpyxl.Workbook) -> dict[str, FxProfile]:
    profiles: dict[str, FxProfile] = {}
    for r in _rows_by_header(wb["演出预设"], "预设ID"):
        pid = _to_str(r.get("预设ID"))
        profiles[pid] = FxProfile(
            profile_id=pid,
            color_role=_to_str(r.get("色彩角色")),
            color_hex=_to_str(r.get("主色HEX")),
            bright_hex=_to_str(r.get("亮色HEX")),
            particle_char=_to_str(r.get("主粒子")),
            warn_cells_high=_to_int(r.get("预警高档字符格")),
            warn_cells_low=_to_int(r.get("预警低档字符格")),
            cast_cells_high=_to_int(r.get("释放高档字符格")),
            cast_cells_low=_to_int(r.get("释放低档字符格")),
            entry_seconds=_to_float(r.get("入场秒")),
            cast_seconds=_to_float(r.get("释放秒")),
            tail_seconds=_to_float(r.get("收尾秒")),
            range_hint=_to_str(r.get("范围提示与演员保护")),
        )
    return profiles


_TUNING_NUMERIC_KEYS = {
    "浅层伤害倍率": ("band_multiplier_1", float),
    "中层伤害倍率": ("band_multiplier_2", float),
    "深层伤害倍率": ("band_multiplier_3", float),
    "边界伤害倍率": ("band_multiplier_4", float),
    "普通伤害上限1": ("damage_cap_tier0_band1", int),
    "普通伤害上限2": ("damage_cap_tier0_band2", int),
    "普通伤害上限3": ("damage_cap_tier0_band3", int),
    "普通伤害上限4": ("damage_cap_tier0_band4", int),
    "强烈伤害上限1": ("damage_cap_tier1_band1", int),
    "强烈伤害上限2": ("damage_cap_tier1_band2", int),
    "强烈伤害上限3": ("damage_cap_tier1_band3", int),
    "强烈伤害上限4": ("damage_cap_tier1_band4", int),
    "核心伤害上限1": ("damage_cap_tier2_band1", int),
    "核心伤害上限2": ("damage_cap_tier2_band2", int),
    "核心伤害上限3": ("damage_cap_tier2_band3", int),
    "核心伤害上限4": ("damage_cap_tier2_band4", int),
    "普通共享起手间隔": ("shared_gap_tier0", int),
    "强烈共享起手间隔": ("shared_gap_tier1", int),
    "核心共享起手间隔": ("shared_gap_tier2", int),
    "普通首次警戒延迟": ("initial_delay_tier0", int),
    "强烈首次警戒延迟": ("initial_delay_tier1", int),
    "核心首次警戒延迟": ("initial_delay_tier2", int),
    "场每拍伤害": ("field_tick_damage", int),
    "场默认寿命": ("field_default_rounds", int),
    "控制保护窗口": ("control_protection_rounds", int),
    "普通/强烈并发预警上限": ("concurrent_warning_limit_ordinary", int),
    "核心房并发预警上限": ("concurrent_warning_limit_core_room", int),
    "召唤全遭遇同时上限": ("summon_limit", int),
    "普通/强烈累计治疗比例": ("heal_fraction_tier01", float),
    "核心累计治疗比例": ("heal_fraction_tier2", float),
    "高档全地图字符预算": ("fx_budget_high", int),
    "低档全地图字符预算": ("fx_budget_low", int),
    "新技能可见格装饰比例": ("fx_decor_ratio", float),
    "高档活跃释放": ("fx_fps_high", int),
    "低档/预警/环境": ("fx_fps_low", int),
    "软控制总移步加值上限": ("soft_control_move_cap", float),
    "普通遭遇中位耗时增幅上限": ("probe_time_gain_ordinary", float),
    "强烈遭遇中位耗时增幅上限": ("probe_time_gain_strong", float),
    "核心遭遇总耗时增幅上限": ("probe_time_gain_core", float),
    "普通/强烈20回合总创伤上限": ("probe_trauma_gain_ordinary", float),
    "核心20回合总创伤上限": ("probe_trauma_gain_core", float),
}


def _parse_tuning(wb: openpyxl.Workbook) -> Mapping[str, Any]:
    values: dict[str, Any] = {}
    raw_notes: dict[str, str] = {}
    for r in _rows_by_header(wb["调参规则"], "参数/来源"):
        label = _to_str(r.get("参数/来源"))
        spec = _TUNING_NUMERIC_KEYS.get(label)
        if spec is None:
            continue
        key, cast = spec
        values[key] = cast(_to_float(r.get("输入值或路径")))
        raw_notes[key] = _to_str(r.get("口径与用途"))
    values["_notes"] = MappingProxyType(raw_notes)
    return MappingProxyType(values)


# ---------------------------------------------------------------------------
# 结构化绑定覆写（与 xlsx「绑定覆写/阶段」文本对应的机器可读版本）
# ---------------------------------------------------------------------------

# (异象ID, 技能ID) → 生效参数覆写。键必须能在《异象配置》原文中找到依据。
_BINDING_OVERRIDES: dict[tuple[str, str], dict[str, Any]] = {
    # 未完成句号：space_fold 按强烈异象参数（核心房不享受核心伤害上限）。
    ("anomaly.unfinished_period", "space_fold"): {"effective_tier": 1},
    # 未完成句号：core_summon 仅生成专属句读残影，不召正式折页人。
    ("anomaly.unfinished_period", "core_summon"): {"summon_variant": "period_comma_shadow"},
    # 空心的灯塔：召唤专属光影，不召正式纸页犬。
    ("anomaly.hollow_lighthouse", "core_summon"): {"summon_variant": "lighthouse_light_shadow"},
    # 沉钟／全声合唱借用初啼声浪，但必须使用弱化参数。
    ("anomaly.sunken_bell", "core_cry"): {
        "shape": "固定曼哈顿半径1",
        "max_cells": 5,
        "direct_base": 2,
        "alert_allies": False,
    },
    ("anomaly.whole_choir", "core_cry"): {
        "shape": "固定曼哈顿半径1",
        "max_cells": 5,
        "direct_base": 2,
        "alert_allies": False,
    },
    # 梦的账房／遗忘档案员借用遗忘批注，但去掉即时创伤，仅想象 −1×2。
    ("anomaly.dream_accountant", "oblivion_mark"): {
        "direct_base": 0,
        "imagination_drain": 2,
    },
    ("anomaly.oblivion_archivist", "oblivion_mark"): {
        "direct_base": 0,
        "imagination_drain": 2,
        "cancel_dispel_on_hit": True,
    },
    # 教学覆写（门扉守卫 1–3 层）：推移无伤害且推向安全格；护折只减 1、最多一次。
    ("anomaly.door_warden", "model_push"): {
        "direct_base": 0,
        "push_to_safe_cell": True,
    },
    ("anomaly.door_warden", "fold_guard"): {"guard_reduce": 1, "guard_max_blocks": 1},
    # 原初冲动：分节共用本体技能状态与冷却，子节不可独立施法。
    ("anomaly.primordial_impulse", "return_charge"): {"segments_share_state": True},
    ("anomaly.primordial_impulse", "model_push"): {"segments_share_state": True},
    # 双重曝光／巡回放映队：被动分身与主动残影共用同一召唤预算。
    ("anomaly.double_exposure", "split_echo"): {"shared_summon_budget": True},
    ("anomaly.projection_crew", "split_echo"): {
        "shared_summon_budget": True,
        "same_source_summon_max": 1,
    },
    # 执笔者：共同校订每阶段一次。
    ("anomaly.scribe", "scribe_revision"): {"once_per_phase": True},
}

# 核心执念阶段解锁（对应设计文档 8.1 与《异象配置》阶段文本）。
_CORE_PHASES: dict[str, dict[str, int]] = {
    "anomaly.hollow_lighthouse": {
        "core_beacon": 1, "choir_guard": 1, "cool_drain": 2, "core_summon": 3,
    },
    "anomaly.unfinished_period": {
        "core_period": 1, "space_fold": 1, "core_summon": 2, "core_pause": 3,
    },
    "anomaly.first_cry": {
        "core_cry": 1, "core_chill": 1, "core_summon": 2, "core_pause": 3,
    },
    "anomaly.scribe": {
        "scribe_stroke": 1, "scribe_erase": 1, "scribe_draft": 2, "scribe_revision": 1,
    },
}


def _tier_from_text(tier_name: str) -> int | None:
    """从「普通异象／强烈异象／核心执念」等文本解析阶级序号。"""
    for name, tier in _TIER_NAME_TO_INT.items():
        if tier_name.startswith(name):
            return tier
    return None


def _parse_bindings(
    wb: openpyxl.Workbook,
    definitions: Mapping[str, SkillDef],
) -> tuple[dict[str, tuple[SkillBinding, ...]], dict[str, AnomalySkillConfig]]:
    bindings: dict[str, list[SkillBinding]] = {}
    configs: dict[str, AnomalySkillConfig] = {}
    for r in _rows_by_header(wb["异象配置"], "异象ID"):
        anomaly_id = _to_str(r.get("异象ID"))
        tier = _tier_from_text(_to_str(r.get("所属阶级")))
        if tier is None:
            raise RuntimeError(
                f"异象 {anomaly_id} 所属阶级非法：{_to_str(r.get('所属阶级'))!r}"
            )
        slots = [
            _to_str(r.get(f"主动{i}"))
            for i in range(1, 5)
            if _to_str(r.get(f"主动{i}"))
        ]
        note = _to_str(r.get("绑定覆写/阶段"))
        phase_map = _CORE_PHASES.get(anomaly_id, {})
        bound: list[SkillBinding] = []
        for slot, skill_id in enumerate(slots):
            if skill_id not in definitions:
                raise RuntimeError(f"异象 {anomaly_id} 绑定未知技能 ID：{skill_id!r}")
            overrides = dict(_BINDING_OVERRIDES.get((anomaly_id, skill_id), {}))
            bound.append(SkillBinding(
                anomaly_id=anomaly_id,
                skill_id=skill_id,
                slot=slot,
                priority=slot,
                phase=phase_map.get(skill_id, 1),
                overrides=overrides,
                override_note=note,
            ))
        bindings[anomaly_id] = tuple(bound)
        configs[anomaly_id] = AnomalySkillConfig(
            anomaly_id=anomaly_id,
            name_zh=_to_str(r.get("现有中文名")),
            tier=tier,
            layer_start=_to_int(r.get("主层起")),
            layer_end=_to_int(r.get("主层止")),
            initial_delay_rounds=_to_int(r.get("初次延迟")),
            global_gap_rounds=_to_int(r.get("共享间隔")),
            override_note=note,
            migration_note=_to_str(r.get("旧机制迁移")),
            status=_to_str(r.get("状态")),
        )
    return bindings, configs


# ---------------------------------------------------------------------------
# 模块级构建（导入时一次性完成）
# ---------------------------------------------------------------------------

def _build() -> tuple[
    dict[str, SkillDef],
    dict[str, tuple[SkillBinding, ...]],
    dict[str, FxProfile],
    dict[str, AnomalySkillConfig],
    Mapping[str, Any],
]:
    if _PREBUILT_PATH.exists():
        return _load_prebuilt()
    wb = _load_workbook()
    try:
        definitions = _parse_skills(wb)
        fx = _parse_fx_profiles(wb)
        tuning = _parse_tuning(wb)
        bindings, configs = _parse_bindings(wb, definitions)
    finally:
        wb.close()
    return definitions, bindings, fx, configs, tuning


# ---------------------------------------------------------------------------
# 预转换 JSON：xlsx 派生数据（技能/绑定/演出/配置/调参五份）的离线快照。
# 生成工具 tools/export_anomaly_skills_json.py；格式只增不改（version 字段）。
# ---------------------------------------------------------------------------

_PREBUILT_VERSION = 1


def _parse_workbook() -> tuple[
    dict[str, SkillDef],
    dict[str, tuple[SkillBinding, ...]],
    dict[str, FxProfile],
    dict[str, AnomalySkillConfig],
    Mapping[str, Any],
]:
    wb = _load_workbook()
    try:
        definitions = _parse_skills(wb)
        fx = _parse_fx_profiles(wb)
        tuning = _parse_tuning(wb)
        bindings, configs = _parse_bindings(wb, definitions)
    finally:
        wb.close()
    return definitions, bindings, fx, configs, tuning


def dump_prebuilt(path: Path | None = None) -> Path:
    """把当前 xlsx 解析结果写成预转换 JSON（构建期工具调用，勿在运行期用）。"""
    from dataclasses import asdict

    definitions, bindings, fx, configs, tuning = _parse_workbook()
    target = path or _PREBUILT_PATH
    payload = {
        "version": _PREBUILT_VERSION,
        "skills": {sid: asdict(d) for sid, d in definitions.items()},
        "bindings": {
            aid: [
                {
                    "anomaly_id": b.anomaly_id, "skill_id": b.skill_id,
                    "slot": b.slot, "priority": b.priority, "phase": b.phase,
                    "overrides": dict(b.overrides), "override_note": b.override_note,
                }
                for b in slots
            ]
            for aid, slots in bindings.items()
        },
        "fx_profiles": {pid: asdict(p) for pid, p in fx.items()},
        "configs": {aid: asdict(c) for aid, c in configs.items()},
        "tuning": {k: v for k, v in tuning.items() if k != "_notes"},
        "tuning_notes": dict(tuning["_notes"]),
    }
    target.write_text(
        json.dumps(payload, ensure_ascii=False, separators=(",", ":")),
        encoding="utf-8")
    return target


def _load_prebuilt() -> tuple[
    dict[str, SkillDef],
    dict[str, tuple[SkillBinding, ...]],
    dict[str, FxProfile],
    dict[str, AnomalySkillConfig],
    Mapping[str, Any],
]:
    try:
        payload = json.loads(_PREBUILT_PATH.read_text(encoding="utf-8"))
        if payload.get("version") != _PREBUILT_VERSION:
            raise RuntimeError(f"预转换版本不符：{payload.get('version')}")
        definitions = {
            sid: SkillDef(**{**d, "damage_by_band": tuple(d["damage_by_band"])})
            for sid, d in payload["skills"].items()
        }
        bindings = {
            aid: tuple(
                SkillBinding(
                    anomaly_id=b["anomaly_id"], skill_id=b["skill_id"],
                    slot=b["slot"], priority=b["priority"], phase=b["phase"],
                    overrides=b["overrides"], override_note=b["override_note"])
                for b in slots)
            for aid, slots in payload["bindings"].items()
        }
        fx = {pid: FxProfile(**p) for pid, p in payload["fx_profiles"].items()}
        configs = {aid: AnomalySkillConfig(**c) for aid, c in payload["configs"].items()}
        values = dict(payload["tuning"])
        values["_notes"] = MappingProxyType(dict(payload["tuning_notes"]))
        tuning = MappingProxyType(values)
        return definitions, bindings, fx, configs, tuning
    except Exception as error:
        raise RuntimeError(f"预转换 JSON 解析失败（{error}）；请重跑 "
                           f"tools/export_anomaly_skills_json.py 或删除该文件回退 xlsx") from error


SKILL_DEFINITIONS, ANOMALY_SKILL_BINDINGS, FX_PROFILES, ANOMALY_SKILL_CONFIGS, TUNING = _build()
from .skill_prose import SKILL_PROSE
for _skill_id, (_effect, _condition, _effect_en, _condition_en) in SKILL_PROSE.items():
    SKILL_DEFINITIONS[_skill_id] = replace(
        SKILL_DEFINITIONS[_skill_id], effect_text=_effect, condition_text=_condition)

from .hospital import HOSPITAL_SKILLS, HOSPITAL_BINDINGS, HOSPITAL_ANOMALIES, HOSPITAL_SKILL_EFFECTS
for key, name, tier, cd, windup, reach, cells, power, channel, recovery, fx in HOSPITAL_SKILLS:
    sid = 'hospital.'+key
    SKILL_DEFINITIONS[sid] = SkillDef(
        sid, name, '医疗意象', tier, cd, windup, reach, '锁定局部范围', cells,
        power, (power,)*4, channel, 3 if key in ('membrane_guard', 'scan_line') else 2 if key == 'filament' else 0, recovery, fx,
        HOSPITAL_SKILL_EFFECTS[key], '目标在视线内；储量不足时停用。',
        '横移、遮挡、利用有限材料与双向净化。', '已实装', fixed_damage=True)
for aid, skills in HOSPITAL_BINDINGS.items():
    tier = ('普通异象', '强烈异象', '核心执念').index(HOSPITAL_ANOMALIES[aid]['anomaly_kind'])
    ANOMALY_SKILL_BINDINGS[aid] = tuple(
        SkillBinding(aid, sid, i, i, 2 if sid == 'hospital.reopen' else 1)
        for i, sid in enumerate(skills))
    ANOMALY_SKILL_CONFIGS[aid] = AnomalySkillConfig(aid, HOSPITAL_ANOMALIES[aid]['name'], tier,
        21, 26, 1, 3 if tier < 2 else 4, '第三区成品值', '专属位面', '已实装')


def band_index_for_layer(layer: int) -> int:
    """主层号 → 层段索引（0–3）。"""
    for i, (lo, hi) in enumerate(_LAYER_BAND_RANGES):
        if lo <= layer <= hi:
            return i
    raise ValueError(f"主层号超出 1–30：{layer}")


def damage_cap(tier: int, band: int) -> int:
    """所属阶级 × 层段 → 单次直接创伤上限。"""
    key = f"damage_cap_tier{tier}_band{band + 1}"
    return int(TUNING[key])


def band_multiplier(band: int) -> float:
    return float(TUNING[f"band_multiplier_{band + 1}"])


def direct_damage(skill_id: str, layer: int, anomaly_id: str | None = None) -> int:
    """一次直接创伤：min(ceil(基础×层段倍率), 施法者阶级上限)。

    ``anomaly_id`` 提供时应用该异象绑定的覆写（含弱化与生效阶级）。
    """
    skill = effective_skill_def(skill_id, anomaly_id) if anomaly_id else SKILL_DEFINITIONS[skill_id]
    if skill.direct_base <= 0:
        return 0
    if skill.fixed_damage:
        return skill.direct_base
    band = band_index_for_layer(layer)
    caster_tier = _caster_tier(skill_id, anomaly_id)
    return min(math.ceil(skill.direct_base * band_multiplier(band)), damage_cap(caster_tier, band))


def _caster_tier(skill_id: str, anomaly_id: str | None) -> int:
    if anomaly_id:
        binding = _binding_for(anomaly_id, skill_id)
        if binding is not None:
            return binding.effective_tier
        return ANOMALY_SKILL_CONFIGS[anomaly_id].tier
    return SKILL_DEFINITIONS[skill_id].default_tier


def _binding_for(anomaly_id: str, skill_id: str) -> SkillBinding | None:
    for binding in ANOMALY_SKILL_BINDINGS.get(anomaly_id, ()):
        if binding.skill_id == skill_id:
            return binding
    return None


def bindings_for(anomaly_id: str) -> tuple[SkillBinding, ...]:
    """某异象的全部主动技能绑定（按配置槽顺序）。"""
    return ANOMALY_SKILL_BINDINGS.get(anomaly_id, ())


def skill_summaries_for(anomaly_id: str, layer: int) -> list[dict[str, Any]]:
    """某异象主动技能的观测摘要（玩家情报展示用）。

    每项：技能 id/名、**完整效果/施放条件/玩家反制**三段文本（图鉴详情与
    观测展开共用）、创伤渠道与**当前层的单次直接创伤值**（绑定覆写——
    弱化、生效阶级——已在 effective_skill_def / direct_damage 里生效）、
    冷却/蓄力/持续/恢复回合与施法距离。无绑定的异象返回空表，调用方
    据此跳过技能行。2026-10-10 补全：文本三段 + 四项参数，观察异象
    一屏看全（不用再翻开发表）。
    """
    summaries: list[dict[str, Any]] = []
    for binding in bindings_for(anomaly_id):
        skill = effective_skill_def(binding.skill_id, anomaly_id)
        summaries.append({
            "skill_id": binding.skill_id,
            "name": skill.name_zh,
            "effect_text": skill.effect_text,
            "condition_text": skill.condition_text,
            "counter_text": skill.counter_text,
            "trauma_channel": skill.trauma_channel,
            "trauma_value": direct_damage(binding.skill_id, layer, anomaly_id),
            "cooldown_rounds": skill.cooldown_rounds,
            "windup_rounds": skill.windup_rounds,
            "duration_rounds": skill.duration_rounds,
            "recovery_rounds": skill.recovery_rounds,
            "max_range": skill.max_range,
        })
    return summaries


def effective_skill_def(skill_id: str, anomaly_id: str | None = None) -> SkillDef:
    """基础定义合并绑定覆写后的生效定义（不修改全局定义）。"""
    base = SKILL_DEFINITIONS[skill_id]
    if anomaly_id is None:
        return base
    binding = _binding_for(anomaly_id, skill_id)
    if binding is None or not binding.overrides:
        return base
    changes = {
        k: v for k, v in binding.overrides.items()
        if k in SkillDef.__dataclass_fields__ and k != "skill_id"
    }
    if not changes:
        return base
    if "direct_base" in changes and "damage_by_band" not in changes:
        # 基础创伤被覆写（弱化/教学）时，四层段按生效阶级重算，
        # 不能沿用默认阶级的预计算结果。
        new_base = int(changes["direct_base"])
        tier = binding.effective_tier
        changes["damage_by_band"] = tuple(
            0 if new_base <= 0 else min(
                math.ceil(new_base * band_multiplier(band)), damage_cap(tier, band)
            )
            for band in range(4)
        )
    return replace(base, **changes)


# ---------------------------------------------------------------------------
# 校验
# ---------------------------------------------------------------------------


# 藏书宫十五项专属技能的玩家文案：逐项对照 core/library_skills.py 的执行逻辑
# 另写（锁定预告格/有限资源与实体遮挡的共用模板已废）。三元组 = 效果／施放条件／反制提示。
LIBRARY_SKILL_TEXTS = {
    'question': ('沿直线追问至多四格，命中一次心智创伤5；起笔之后落点固定，不再跟着你转向。',
                 '目标在视线内，且两人之间有可以落笔的直线。', '横移出线，或用实体遮挡。'),
    'threshold': ('把两格内一扇还开着的门带上，两回合后它自己重新敞开。',
                  '两格内有一扇空置的门。', '门只在侧廊上；等它自己敞开。'),
    'ink': ('把你此刻站着的格子泼成墨迹，停留三回合；墨里行走会拖沓。',
            '目标可见。', '别停在原地；墨迹过几拍就干了。'),
    'omit': ('把自己身上的一项负面征兆整个省略掉，像删去一个多余的词。',
             '它身上确有负面征兆。', '省略掉的就是没有了。'),
    'return': ('翻回它上一处站过的位置，像书签跳回上一回目。',
               '上一处位置在三格以内，且仍然空着。', '占住它回来的位置。'),
    'lamp': ('把四格内一盏阅读灯的光借走两回合；光跟着它走，原处暗下去。',
             '四格内有一盏阅读灯。', '灯最终会归还原位。'),
    'foldstep': ('折过一页书：挪进三格内一扇挨着书架的空格。',
                 '三格内有一格空地挨着书架。', '把书架旁的空地占住。'),
    'bind': ('收走身旁一叠干纸替自己压上纸护，两回合内你打它会更轻；纸堆会薄下去。',
             '身旁有还没烧完的干纸堆，且它身上还没有纸护。', '抢先烧掉纸堆，或等纸护耗尽。'),
    'stroke': ('沿直线划出至多六格的长线，命中一次心智创伤6。',
               '目标在视线内，且两人之间有可以落笔的直线。', '横移出线，或用实体遮挡。'),
    'revision': ('替四格内的一名同伴校订一遍，压上两回合的纸护。这场里它只做一次。',
                 '四格内有还没拿到纸护的同伴。', '同伴的纸护只能挡一下。'),
    'echo': ('把自己誊出几个互相出入的版本：身旁空格里最多多出两份单薄的誊影，各撑八回合。',
             '身旁有空格，且连同誊影在内还不足两份。', '誊影只有一层纸厚。'),
    'blank': ('把身旁至多两格抹成空白，停留两回合；空白里落不下别的把戏。',
              '身旁有空格。', '空白不挡路，只是什么都留不住。'),
    'index_line': ('沿直线归出至多七格的目录长线，命中一次心智创伤7；被线穿过的目标眼前会发花。',
                   '目标在视线内，且两人之间有可以落笔的直线。', '横移出线，或用实体遮挡。'),
    'delete': ('把你周围的格子涂掉至多两格，墨迹停留三回合；正中的主廊不涂。',
               '目标可见。', '主廊上没有墨。'),
    'last_page': ('落下至多七格的最后一页：先写前半行，等你行动之后，后半行跟着落下；每半行各带一次心智创伤7。',
                  '目标在视线内，且两人之间有可以落笔的直线。', '后半行只认预告的格子。'),
}
from .library import LIBRARY_SKILLS, LIBRARY_BINDINGS, LIBRARY_ANOMALIES
for key, name, tier, cd, windup, reach, cells, power, duration in LIBRARY_SKILLS:
    sid = 'library.'+key
    effect, condition, counter = LIBRARY_SKILL_TEXTS[key]
    SKILL_DEFINITIONS[sid] = SkillDef(sid,name,'文学意象',tier,cd,windup,reach,'锁定局部范围',cells,
        power,(power,)*4,'心智' if power else '无',duration,1,'needle' if power else 'field',
        effect,condition,counter,'已实装',fixed_damage=True)
for aid, skills in LIBRARY_BINDINGS.items():
    tier=('普通异象','强烈异象','核心执念').index(LIBRARY_ANOMALIES[aid]['anomaly_kind'])
    ANOMALY_SKILL_BINDINGS[aid]=tuple(SkillBinding(aid,sid,i,i,2 if sid in ('library.delete','library.last_page') else 1) for i,sid in enumerate(skills))
    ANOMALY_SKILL_CONFIGS[aid]=AnomalySkillConfig(aid,LIBRARY_ANOMALIES[aid]['name'],tier,16,22,1,2 if tier==2 else 3,'专属成品值','专属位面','已实装')

# 漏字者的「吞音」（规格 §7.3）：Python 定义，不进工作簿。创伤渠道记「无」、
# 基础创伤记 0——它造成的是**想象创伤**，由 core/tamper_risks.py 的解析器直接
# 结算，绕开层段倍率与阶级上限（那是主梦直接创伤的口径）。
from .tamper_anomalies import TAMPER_ANOMALIES, TAMPER_BINDINGS, TAMPER_SKILLS, OMISSION_SKILL_TEXTS
for key, name, tier, cd, windup, reach, cells, power, duration in TAMPER_SKILLS:
    sid = 'tamper.'+key
    effect, condition, counter = OMISSION_SKILL_TEXTS[key]
    SKILL_DEFINITIONS[sid] = SkillDef(sid, name, '篡改风险', tier, cd, windup, reach,
        '单格锁定', cells, power, (power,)*4, '无', duration, 0, 'sound',
        effect, condition, counter, '已实装')
for aid, skills in TAMPER_BINDINGS.items():
    tier = ('普通异象','强烈异象','核心执念').index(TAMPER_ANOMALIES[aid]['anomaly_kind'])
    ANOMALY_SKILL_BINDINGS[aid] = tuple(
        SkillBinding(aid, sid, i, i, 1) for i, sid in enumerate(skills))
    ANOMALY_SKILL_CONFIGS[aid] = AnomalySkillConfig(aid, TAMPER_ANOMALIES[aid]['name'], tier,
        1, 30, 1, 3, '篡改风险专属', '不进投放池', '已实装')

# ---- 混乱极高专属异象（2026-10-27 玩家需求）----
# 玩家要求：专属异象「比较强、有强力技能」。六技能基础创伤 8–11，按强烈
# 异象阶级上限结算（深层顶到 6/7 点）；两个控制项（失焦征兆、迟帧光场）
# 无直接创伤。施法/结算复用 EnemySkillMixin 既有 handler，只有两个范围技能
# 与吸血/掏空在 core/chaos_skills.py 里补少量专属 handler（不重写引擎）。
# 位置必须在 hospital/library/tamper 三批专属注册之后、2026-10-15 混乱化
# 调参 pass 之前，这样新技能的射程/格数/持续也走同一套统一口径。
from .chaos_anomalies import (
    CHAOS_ANOMALIES,
    CHAOS_BINDINGS,
    CHAOS_SKILLS,
    CHAOS_SKILL_META,
    CHAOS_SKILL_TEXTS,
)

for key, name, tier, cd, windup, reach, cells, power, duration in CHAOS_SKILLS:
    sid = 'chaos.' + key
    effect, condition, counter = CHAOS_SKILL_TEXTS[key]
    shape, channel, fx_profile, recovery = CHAOS_SKILL_META[key]
    damage_by_band = tuple(
        0 if power <= 0 else min(math.ceil(power * band_multiplier(b)), damage_cap(tier, b))
        for b in range(4)
    )
    SKILL_DEFINITIONS[sid] = SkillDef(
        sid, name, '混乱意象', tier, cd, windup, reach, shape, cells, power,
        damage_by_band, channel, duration, recovery, fx_profile,
        effect, condition, counter, '已实装')
for aid, skills in CHAOS_BINDINGS.items():
    tier = ('普通异象', '强烈异象', '核心执念').index(CHAOS_ANOMALIES[aid]['anomaly_kind'])
    ANOMALY_SKILL_BINDINGS[aid] = tuple(
        SkillBinding(aid, sid, i, i, 1) for i, sid in enumerate(skills))
    ANOMALY_SKILL_CONFIGS[aid] = AnomalySkillConfig(
        aid, CHAOS_ANOMALIES[aid]['name'], tier, 1, 27, 1, 3,
        '混乱极高专属', '不进投放池', '已实装')

# ---------------------------------------------------------------------------
# 2026-10-15 混乱化改造调参 pass：异象技能全面放开射程、格数与持续时间。
# 「很多技能实际上是视野内或无限范围，但都被规定了小范围」——统一放开：
# - 施法距离：全部抬到切比雪夫 8（视野级），视野内就该打得着；
# - 射线/范围格数上限 max_cells ≥13：射程放开了，弹道也得够长；
# - 场/征兆持续：非零 duration_rounds 一律 +2 回合，改造的场面留得更久。
# 伤害（direct_base/damage_by_band）与 CD 一律不动——「破坏力」提现在场面。
# 位置：必须在 hospital/library/tamper 三批专属注册全部完成之后（模块底部）。
for _sid, _def in list(SKILL_DEFINITIONS.items()):
    _kw: dict[str, Any] = {}
    if _def.max_range < 8:
        _kw["max_range"] = 8
    if _def.max_cells < 13:
        _kw["max_cells"] = 13
    if _def.duration_rounds:
        _kw["duration_rounds"] = _def.duration_rounds + 2
    if _kw:
        SKILL_DEFINITIONS[_sid] = replace(_def, **_kw)

_EXPECTED_SKILL_COUNT = (43 + len(HOSPITAL_SKILLS) + len(LIBRARY_SKILLS)
                        + len(TAMPER_SKILLS) + len(CHAOS_SKILLS))
_EXPECTED_ANOMALY_COUNT = (56 + len(HOSPITAL_BINDINGS) + len(LIBRARY_BINDINGS)
                           + len(TAMPER_BINDINGS) + len(CHAOS_BINDINGS))
_EXPECTED_BINDING_COUNT = (87 + sum(map(len, HOSPITAL_BINDINGS.values()))
                           + sum(map(len, LIBRARY_BINDINGS.values()))
                           + sum(map(len, TAMPER_BINDINGS.values()))
                           + sum(map(len, CHAOS_BINDINGS.values())))
# 槽位数量阶级规则：普通 0–2 / 强烈 2–3 / 核心 4。
_SLOT_LIMITS = ((0, 2), (2, 3), (4, 4))


def _anomaly_tier_from_definitions(anomaly_id: str) -> int | None:
    from .anomalies import ANOMALY_DEFINITIONS

    definition = ANOMALY_DEFINITIONS.get(anomaly_id)
    if definition is None:
        return None
    return _tier_from_text(_to_str(definition.get("anomaly_kind")))


def validate() -> list[str]:
    """返回全部数据一致性问题；空列表表示通过。"""
    errors: list[str] = []

    if len(SKILL_DEFINITIONS) != _EXPECTED_SKILL_COUNT:
        errors.append(f"技能定义数量 {len(SKILL_DEFINITIONS)} != {_EXPECTED_SKILL_COUNT}")
    if len(ANOMALY_SKILL_BINDINGS) != _EXPECTED_ANOMALY_COUNT:
        errors.append(
            f"异象绑定数量 {len(ANOMALY_SKILL_BINDINGS)} != {_EXPECTED_ANOMALY_COUNT}"
        )
    total_slots = sum(len(b) for b in ANOMALY_SKILL_BINDINGS.values())
    if total_slots != _EXPECTED_BINDING_COUNT:
        errors.append(f"绑定槽位总数 {total_slots} != {_EXPECTED_BINDING_COUNT}")

    # 每个定义都有使用者；每个绑定引用已知异象。
    used_skills = {
        b.skill_id for bindings in ANOMALY_SKILL_BINDINGS.values() for b in bindings
    }
    for skill_id in sorted(set(SKILL_DEFINITIONS) - used_skills):
        errors.append(f"孤立技能定义（无绑定使用者）：{skill_id}")

    from .anomalies import ANOMALY_DEFINITIONS

    for anomaly_id, config in ANOMALY_SKILL_CONFIGS.items():
        if anomaly_id not in ANOMALY_DEFINITIONS:
            errors.append(f"绑定引用了未知异象 ID：{anomaly_id}")
            continue
        def_tier = _anomaly_tier_from_definitions(anomaly_id)
        if def_tier is None:
            errors.append(f"异象 {anomaly_id} 缺少 anomaly_kind，无法判定阶级")
        elif def_tier != config.tier:
            errors.append(
                f"异象 {anomaly_id} 阶级不一致：xlsx={TIER_NAMES[config.tier]} "
                f"定义={TIER_NAMES[def_tier]}"
            )
        lo, hi = _SLOT_LIMITS[config.tier]
        if config.tier == 1 and ANOMALY_DEFINITIONS[anomaly_id].get('spawn_realms'):
            lo, hi = 1, 2
        count = len(ANOMALY_SKILL_BINDINGS[anomaly_id])
        if not lo <= count <= hi:
            errors.append(
                f"异象 {anomaly_id}（{TIER_NAMES[config.tier]}）绑定槽位 {count} "
                f"超出 [{lo},{hi}]"
            )

    for skill in SKILL_DEFINITIONS.values():
        sid = skill.skill_id
        if skill.cooldown_rounds < 0 or skill.windup_rounds < 0 or skill.recovery_rounds < 0:
            errors.append(f"技能 {sid} CD/蓄力/恢复为负")
        if skill.duration_rounds < 0:
            errors.append(f"技能 {sid} 持续为负")
        if skill.max_cells < 1:
            errors.append(f"技能 {sid} 格数上限 < 1")
        if skill.trauma_channel not in ("心智", "清醒", "无"):
            errors.append(f"技能 {sid} 创伤渠道非法：{skill.trauma_channel!r}")
        if skill.fx_profile not in FX_PROFILES:
            errors.append(f"技能 {sid} 引用未知演出预设：{skill.fx_profile!r}")
        # 伤害上限阶层合法：预计算四层段 = min(ceil(基础×倍率), 默认阶级上限)。
        for band in range(4):
            if skill.fixed_damage:
                if skill.damage_by_band[band] != skill.direct_base or not 0 <= skill.direct_base <= 12:
                    errors.append(f'位面技能成品创伤非法：{sid}')
                continue
            cap = damage_cap(skill.default_tier, band)
            expected = 0 if skill.direct_base == 0 else min(
                math.ceil(skill.direct_base * band_multiplier(band)), cap
            )
            actual = skill.damage_by_band[band]
            if actual != expected:
                errors.append(
                    f"技能 {sid} 层段{band + 1}伤害 {actual} != 公式结果 {expected}（上限 {cap}）"
                )
            if not 0 <= actual <= cap:
                errors.append(f"技能 {sid} 层段{band + 1}伤害 {actual} 超出阶级上限 {cap}")

    # 覆写后的伤害值不得超过生效阶级上限。
    for anomaly_id, bindings in ANOMALY_SKILL_BINDINGS.items():
        tier = _caster_tier_for_validation(anomaly_id)
        for binding in bindings:
            skill = effective_skill_def(binding.skill_id, anomaly_id)
            if skill.direct_base <= 0:
                continue
            if skill.fixed_damage:
                if not ANOMALY_DEFINITIONS[anomaly_id].get('spawn_realms'):
                    errors.append(f'固定创伤技能必须绑定专属位面：{anomaly_id}')
                continue
            for band in range(4):
                cap = damage_cap(tier, band)
                expected = min(math.ceil(skill.direct_base * band_multiplier(band)), cap)
                if skill.damage_by_band[band] not in (expected, 0) and skill.damage_by_band[band] > cap:
                    errors.append(
                        f"覆写后技能 {binding.skill_id}@{anomaly_id} 层段{band + 1}伤害 "
                        f"{skill.damage_by_band[band]} 超出生效阶级上限 {cap}"
                    )

    return errors


def _caster_tier_for_validation(anomaly_id: str) -> int:
    config = ANOMALY_SKILL_CONFIGS[anomaly_id]
    tier = config.tier
    for binding in ANOMALY_SKILL_BINDINGS[anomaly_id]:
        if "effective_tier" in binding.overrides:
            tier = int(binding.overrides["effective_tier"])
            break
    return tier


def validate_or_raise() -> None:
    errors = validate()
    if errors:
        raise RuntimeError("异象主动技能数据校验失败：\n" + "\n".join(f"- {e}" for e in errors))


validate_or_raise()


# ---------------------------------------------------------------------------
# 最小接入：向 anomalies.py 定义追加 active_skill_bindings 字段
# ---------------------------------------------------------------------------

def attach_active_skill_bindings() -> None:
    """把绑定写进 ``ANOMALY_DEFINITIONS[aid]["active_skill_bindings"]``。

    只新增字段，不改动旧 ``skill`` / ``skills`` 字段（旧周期结算仍由 core
    使用）；未配置主动技能的异象不设该字段。
    """
    from .anomalies import ANOMALY_DEFINITIONS

    for anomaly_id, bindings in ANOMALY_SKILL_BINDINGS.items():
        if not bindings:
            continue  # 零主动异象不设该字段（普通异象 0–2 个，6 种为 0）
        definition = ANOMALY_DEFINITIONS.get(anomaly_id)
        if definition is not None:
            definition["active_skill_bindings"] = bindings


attach_active_skill_bindings()
